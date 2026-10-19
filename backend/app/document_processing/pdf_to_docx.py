"""PDF có lớp chữ -> DOCX, chuẩn hóa bố cục cơ bản để đọc/phân tích.

Căn cứ: Phụ lục I, Mục 1, Phần I và Phần III.2 của Nghị định
187/2025/NĐ-CP (tệp nguồn ngày 01/07/2025). Chỉ tự động hóa một phần
thể thức; không chứng nhận toàn bộ văn bản tuân thủ pháp luật.
Không sửa ký tự ngoài khoảng trắng, số liệu hoặc dẫn chiếu.
Chỉ xuất DOCX và report JSON; tệp trung gian tự xóa khi kết thúc.

Cài: python -m pip install pdf2docx python-docx "pymupdf>=1.24.2"
Chạy: python pdf_to_docx.py input.pdf converted.docx
Thử DOCX sẵn có: python pdf_to_docx.py converted.docx cleaned.docx --from-docx
Với S/tài liệu khác: thêm --kind other (không tách Điều/Khoản/Điểm).
Mặc định loại số trang theo bbox trên bản PDF tạm; --keep-page-numbers để giữ.
Sửa khoảng trắng dựa trên chữ/vị trí trong PDF, không cần DOCX đối chiếu.
Không phục hồi bảng đã mất; không cam kết tái tạo DOCX gốc giống hệt.
"""
import argparse
import json
import re
import shutil
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ARTICLE = r"Điều\s*\d+[a-zđĐ]?\s*\."
CLAUSE = r"\d{1,3}\.(?!\d)"
POINT = r"[a-zđ]\)"
LABEL = rf"(?:{ARTICLE}|{CLAUSE}|{POINT})"
CHAPTER = re.compile(r"^(?:(?:Phần|Chương)\s+[IVXLCDM]+|(?:Mục|Tiểu mục)\s+\d+)\s*$")
ARTICLE_START = re.compile(rf"^{ARTICLE}")
CLAUSE_START = re.compile(rf"^{CLAUSE}")
POINT_START = re.compile(rf"^{POINT}")


def compact(text):
    return re.sub(r"\s", "", text)


def all_paragraphs(container):
    """Duyệt paragraph/bảng theo thứ tự thực trong XML, kể cả ô gộp."""
    from docx.table import Table
    root = container._element
    if root.tag == qn('w:document'):
        root = root.body
    seen = set()
    for child in root:
        if child.tag == qn('w:p'):
            yield Paragraph(child, container)
        elif child.tag == qn('w:tbl'):
            for row in Table(child, container).rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        yield from all_paragraphs(cell)


def set_font(run):
    run.font.name = "Times New Roman"
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "Times New Roman")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)


def safe_text_paragraph(p):
    # Không viết lại paragraph có ảnh, field, hyperlink hoặc textbox.
    allowed = {qn("w:r"), qn("w:pPr")}
    run_allowed = {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}
    return all(child.tag in allowed for child in p._p) and all(
        child.tag in run_allowed for r in p.runs for child in r._r
    )


def split_units(text):
    """Chỉ tách nhãn sau xuống dòng/tab, tránh tách dẫn chiếu giữa câu."""
    text = text.strip()
    parts = re.split(rf"[\n\t]+[ \t]*(?={LABEL})", text)
    result = []
    for part in parts:
        # Xuống dòng còn lại coi là xuống dòng trình bày trong cùng đơn vị.
        part = re.sub(r"\s+", " ", part).strip()
        part = re.sub(r"^Điều\s*(\d+[a-zđĐ]?)\s*\.\s*", r"Điều \1. ", part)
        part = re.sub(r"^(\d{1,3})\.(?!\d)\s*", r"\1. ", part)
        part = re.sub(r"^([a-zđ])\)\s*", r"\1) ", part)
        if part:
            result.append(part)
    assert compact(text) == compact("".join(result)), "Chuẩn hóa làm đổi ký tự"
    return result


def replace_with_units(p, units):
    """Dựng các paragraph có định dạng cơ bản; giữ ghi chép thay đổi trong report."""
    parent, anchor = p._parent, p._p
    properties = deepcopy(p._p.pPr) if p._p.pPr is not None else None
    created = []
    for index, text in enumerate(units):
        element = OxmlElement("w:p")
        if properties is not None:
            props = deepcopy(properties)
            # Section break thuộc đoạn cuối, không nhân lên mỗi đơn vị.
            if index != len(units) - 1:
                for section in list(props.findall(qn("w:sectPr"))):
                    props.remove(section)
            element.append(props)
        anchor.addprevious(element)
        new = Paragraph(element, parent)
        new.add_run(text)
        created.append(new)
    anchor.getparent().remove(anchor)
    return created


def format_unit(p, subdivision_title=False):
    text = p.text.strip()
    chapter = bool(CHAPTER.match(text))
    article = bool(ARTICLE_START.match(text))
    uppercase_title = subdivision_title and text.isupper() and any(c.isalpha() for c in text)
    p.style = "Heading 1" if chapter else "Heading 2" if article else "Normal"
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.2
    fmt.left_indent = Cm(1.25) if article else Cm(0)
    fmt.right_indent = Cm(0)
    fmt.first_line_indent = Cm(0) if chapter or article or uppercase_title else Cm(1.25)
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if chapter or uppercase_title else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.keep_with_next = chapter or article or uppercase_title
    fmt.keep_together = chapter or article
    fmt.widow_control = True
    for run in p.runs:
        set_font(run)
        run.font.size = Pt(14)
        run.bold = chapter or article or uppercase_title
        run.italic = False


def inspect_pdf(path):
    import fitz
    numbered_pages = set()
    suspicious = []
    with fitz.open(path) as pdf:
        if pdf.needs_pass:
            raise ValueError("PDF có mật khẩu; hãy cung cấp bản mở được.")
        if not len(pdf):
            raise ValueError("PDF không có trang.")
        for number, page in enumerate(pdf, 1):
            text = page.get_text()
            if not text.strip():
                raise ValueError(f"Trang {number} không có chữ trích xuất được; chưa hỗ trợ OCR.")
            if len(re.sub(r"\s", "", text)) < 30:
                suspicious.append(number)
            # Chỉ nhận số khớp trang vật lý, nằm sát đầu/cuối trang và gần giữa.
            for word in page.get_text("words"):
                x0, y0, x1, y1, token = word[:5]
                if (token == str(number)
                    and (y1 < page.rect.height * .09 or y0 > page.rect.height * .91)
                    and abs((x0 + x1) / 2 - page.rect.width / 2) < page.rect.width * .12):
                    numbered_pages.add(token)
    return numbered_pages, suspicious


def body_paragraphs(doc):
    """Chỉ lấy phần nội dung V1; không suy diễn khối ký hoặc phụ lục."""
    active = False
    result = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if re.match(r"^Nơi nhận\s*:|^Phụ lục(?:\s|$)", t, re.I):
            break
        if ARTICLE_START.match(t) or CHAPTER.match(t):
            active = True
        if active:
            result.append(p)
    return result


class PdfEvidence:
    """Chỉ mục chữ và khoảng trắng từ PDF, không dùng DOCX đối chiếu.

    Chỉ sửa đoạn khi chuỗi ký tự không tính khoảng trắng khớp duy nhất.
    Không suy đoán từ tiếng Việt hay thay chữ/số/dấu câu.
    """
    def __init__(self, path):
        import fitz
        from statistics import median
        self.chars, self.lines, self.entries = [], [], []
        self.inferred_spaces = 0
        self.page_images = []
        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, 1):
                if page.get_images():
                    self.page_images.append(page_number)
                raw = page.get_text('rawdict', sort=True,
                                    flags=fitz.TEXTFLAGS_RAWDICT & ~fitz.TEXT_PRESERVE_IMAGES)
                for block in raw['blocks']:
                    for line in block.get('lines', []):
                        horizontal = abs(line.get('dir', (1, 0))[0] - 1) < .01
                        spans = line.get('spans', [])
                        spaces = [c['bbox'][2] - c['bbox'][0] for span in spans
                                  for c in span['chars'] if c['c'].isspace()
                                  and c['bbox'][2] > c['bbox'][0]]
                        space_width = median(spaces) if spaces else None
                        line_id = len(self.lines)
                        record = {'page': page_number, 'bbox': list(line['bbox']),
                                  'page_width': page.rect.width, 'page_height': page.rect.height,
                                  'start': len(self.chars), 'horizontal': horizontal}
                        previous = None
                        whitespace = False
                        for span in spans:
                            size = span.get('size', 12)
                            for char in span['chars']:
                                token = char['c']
                                if token.isspace():
                                    whitespace = True
                                    continue
                                inferred = False
                                if previous and not whitespace and horizontal:
                                    gap = char['bbox'][0] - previous['bbox'][2]
                                    threshold = max(1.2, .8 * space_width if space_width else .2 * size)
                                    if gap > threshold:
                                        whitespace = True
                                        inferred = True
                                        self.inferred_spaces += 1
                                for k, value in enumerate(token):
                                    self.chars.append(value)
                                    self.entries.append({'line': line_id,
                                                         'space': whitespace if k == 0 else False,
                                                         'inferred': inferred if k == 0 else False})
                                whitespace = False
                                previous = char
                        record['end'] = len(self.chars)
                        record['text'] = ''.join(self.chars[record['start']:record['end']])
                        self.lines.append(record)
        self.text = ''.join(self.chars)

    def locate(self, text):
        key = compact(text)
        if len(key) < 12:
            return None  # Nhãn ngắn hoặc chuỗi lặp không đủ để xác minh.
        start = self.text.find(key)
        if start < 0 or self.text.find(key, start + 1) >= 0:
            return None
        end = start + len(key)
        if any(not self.lines[x['line']]['horizontal'] for x in self.entries[start:end]):
            return None
        return start, end

    def location(self, match):
        first = self.lines[self.entries[match[0]]['line']]
        last = self.lines[self.entries[match[1]-1]['line']]
        return {'page_start': first['page'], 'page_end': last['page'],
                'first_line_bbox': first['bbox'], 'last_line_bbox': last['bbox']}

    def repair_spaces(self, text, match):
        tokens = list(re.finditer(r'\S', text))
        result = [text[:tokens[0].start()], tokens[0].group()]
        uncertain = 0
        for offset, token in enumerate(tokens[1:], 1):
            previous = tokens[offset-1]
            original = text[previous.end():token.start()]
            current = self.entries[match[0]+offset]
            left = self.entries[match[0]+offset-1]
            structural = ('\n' in original or '\t' in original) and re.match(LABEL, text[token.start():])
            if structural:
                gap = original
            elif current['line'] == left['line']:
                gap = ' ' if current['space'] else ''
            else:
                # Không đoán ghép/tách từ tại ranh giới dòng chỉ từ chữ.
                gap = original
                if not original and previous.group().isalnum() and token.group().isalnum():
                    uncertain += 1
            result.extend([gap, token.group()])
        result.append(text[tokens[-1].end():])
        return ''.join(result), uncertain

    def continuation(self, left, right):
        """Bằng chứng nối đoạn: liền chữ + qua trang + ở mép thân trang."""
        if left[1] != right[0]:
            return False
        a = self.lines[self.entries[left[1]-1]['line']]
        b = self.lines[self.entries[right[0]]['line']]
        return (b['page'] == a['page'] + 1
                and left[1] == a['end'] and right[0] == b['start']
                and a['bbox'][3] > a['page_height'] * .70
                and b['bbox'][1] < b['page_height'] * .25
                and abs(a['bbox'][0] - b['bbox'][0]) < 20)


def preserve_runs_with_whitespace(p, text):
    """Thay khoảng trắng, giữ run/font nhấn mạnh của các ký tự gốc."""
    if compact(p.text) != compact(text):
        raise RuntimeError('Từ chối sửa: có thay đổi ký tự ngoài khoảng trắng.')
    owners = [i for i, run in enumerate(p.runs) for c in run.text if not c.isspace()]
    if len(owners) != len(compact(text)):
        raise RuntimeError('Không ánh xạ được ký tự vào run.')
    values = [''] * len(p.runs)
    cursor = 0
    for char in text:
        if char.isspace():
            owner = owners[min(cursor, len(owners)-1)] if owners else 0
        else:
            owner = owners[cursor]
            cursor += 1
        values[owner] += char
    for run, value in zip(p.runs, values):
        run.text = value


def repair_from_pdf(doc, evidence, report):
    """Áp dụng trước chuẩn hóa nhãn; kiểm chứng độc lập từng paragraph."""
    records = []
    matches = 0
    for index, p in enumerate(all_paragraphs(doc)):
        if not p.text.strip():
            continue
        match = evidence.locate(p.text)
        if not match or not safe_text_paragraph(p):
            records.append({'paragraph_before_normalization_0based': index,
                            'reason': 'no_unique_pdf_match_or_complex_element', 'excerpt': p.text[:120]})
            continue
        matches += 1
        corrected, uncertain = evidence.repair_spaces(p.text, match)
        if corrected != p.text:
            report['changes'].append({'action': 'pdf_whitespace_repair',
                                      'before': p.text, 'after': corrected,
                                      'pdf_location': evidence.location(match)})
            preserve_runs_with_whitespace(p, corrected)
        if uncertain:
            records.append({'paragraph_before_normalization_0based': index,
                            'reason': 'uncertain_word_boundary_at_pdf_line_break',
                            'count': uncertain, 'excerpt': p.text[:120]})
    report['pdf_alignment'] = {'matched_paragraphs': matches, 'needs_review': records,
                               'geometry_space_candidates': evidence.inferred_spaces,
                               'pages_with_images': evidence.page_images}


def join_pdf_continuations(doc, evidence, report):
    items = [p for p in body_paragraphs(doc) if p.text.strip()]
    i = 0
    while i + 1 < len(items):
        p, following = items[i:i+2]
        a, b = p.text.rstrip(), following.text.lstrip()
        left, right = evidence.locate(a), evidence.locate(b)
        # Không nối tiêu đề/nhãn, danh sách mới, đoạn kết thúc câu hoặc trích dẫn.
        can_join = (left and right and safe_text_paragraph(p) and safe_text_paragraph(following)
                    and not ARTICLE_START.match(a) and not CHAPTER.match(a)
                    and b and b[0].islower() and not re.match(LABEL, b)
                    and a and (a[-1] == ',' or a[-1].isalnum())
                    and evidence.continuation(left, right))
        # Không được vượt qua bảng/ảnh hoặc paragraph có dữ liệu khác.
        node = p._p.getnext()
        while can_join and node is not following._p:
            if node is None or node.tag != qn('w:p'):
                can_join = False
                break
            between = Paragraph(node, p._parent)
            if between.text.strip() or not safe_text_paragraph(between):
                can_join = False
                break
            node = node.getnext()
        if can_join:
            report['changes'].append({'action': 'join_pdf_page_continuation',
                                      'before': [a, b], 'after': a + ' ' + b,
                                      'pdf_location': evidence.location((left[0], right[1]))})
            p.add_run(' ')
            for run in following.runs:
                p._p.append(deepcopy(run._r))
            following.clear()  # Giữ section break; không xóa header/footer.
            items.pop(i+1)
        else:
            i += 1


def tag_document_regions(doc, report):
    """Phân vùng bằng style; giữ nguyên chữ. Parser có thể bỏ qua vùng ngoài thân."""
    from docx.enum.style import WD_STYLE_TYPE
    for name in ('Legal Preamble', 'Legal Closing', 'Legal Appendix'):
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles['Normal']
    region = 'preamble'
    report['regions'] = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if re.match(r'^Phụ lục(?:\s|$)', t, re.I):
            region = 'appendix'
        elif re.match(r'^Nơi nhận\s*:', t):
            region = 'closing'
        elif region == 'preamble' and (ARTICLE_START.match(t) or CHAPTER.match(t)):
            region = 'body'
        if region != 'body' and t:
            p.style = {'preamble': 'Legal Preamble', 'closing': 'Legal Closing',
                       'appendix': 'Legal Appendix'}[region]
        if t:
            report['regions'].append({'output_paragraph_0based': i, 'region': region})
    report['warnings'].append('Phân vùng theo nhãn chỉ là hỗ trợ; chưa khôi phục thứ tự cột ở nơi nhận/chữ ký, không loại bỏ nội dung.')


def remove_pdf_margin_numbers(source, destination):
    """Xóa đúng bbox số trang trên bản PDF tạm, không xóa theo giá trị trong DOCX."""
    import fitz
    removed = []
    with fitz.open(source) as pdf:
        for index, page in enumerate(pdf, 1):
            words = page.get_text('words')
            for word in words:
                x0, y0, x1, y1, token = word[:5]
                if not (token == str(index)
                        and (y1 < page.rect.height * .09 or y0 > page.rect.height * .91)
                        and abs((x0+x1)/2-page.rect.width/2) < page.rect.width*.12):
                    continue
                # Một dòng chỉ có duy nhất số; không đụng số nằm trong câu.
                line = [w for w in words if w[5:7] == word[5:7]]
                if len(line) != 1:
                    continue
                rect = fitz.Rect(x0, y0, x1, y1)
                page.add_redact_annot(rect, fill=False)
                removed.append({'page': index, 'text': token, 'bbox': list(rect)})
            # Không xóa ảnh hoặc đường kẻ. Chỉ sửa bản tạm, giữ PDF gốc.
            page.apply_redactions(images=0, graphics=0, text=0)
        pdf.save(destination)
    return removed


def reflow_text_layout(doc, report):
    """Bỏ ngắt trang/cột do chuyển đổi cho V1 chữ thuần, đồng nhất khổ trang.

    Không áp dụng cho bảng, phụ lục, trang ngang hoặc header/footer có nội dung.
    """
    sections = list(doc.sections)
    if (doc.tables or any(s.page_width > s.page_height for s in sections)
            or any(p.style.name == 'Legal Appendix' for p in doc.paragraphs)
            or any(part.tables or any(p.text.strip() or not safe_text_paragraph(p) for p in part.paragraphs)
                   for s in sections for part in (s.header, s.footer, s.first_page_header,
                       s.first_page_footer, s.even_page_header, s.even_page_footer))):
        report['warnings'].append('Giữ bố cục section vì có bảng/phụ lục/trang ngang/header/footer cần bảo toàn.')
        return
    for p in doc.paragraphs:
        props = p._p.pPr
        if props is not None:
            for sect in list(props.findall(qn('w:sectPr'))):
                props.remove(sect)
        for br in list(p._p.iter(qn('w:br'))):
            if br.get(qn('w:type')) in ('page', 'column'):
                br.getparent().remove(br)
        p.paragraph_format.page_break_before = False
    final = doc.sections[-1]._sectPr
    for name in ('cols', 'type'):
        for child in list(final.findall(qn('w:' + name))):
            final.remove(child)
    cols = OxmlElement('w:cols');cols.set(qn('w:num'), '1');final.append(cols)
    for p in list(doc.paragraphs):
        if not p.text.strip() and safe_text_paragraph(p):
            p._p.getparent().remove(p._p)
            continue
        if p.style.name == 'Legal Preamble':
            units = [x.strip() for x in re.split(r'[\t\n]+', p.text) if x.strip()]
            paras = replace_with_units(p, units) if safe_text_paragraph(p) and len(units)>1 else [p]
            for item in paras:
                item.style = 'Legal Preamble'
                fmt = item.paragraph_format
                fmt.left_indent = fmt.right_indent = Cm(0)
                fmt.first_line_indent = Cm(1.25) if item.text.startswith('Căn cứ ') else Cm(0)
                fmt.space_before = Pt(0);fmt.space_after = Pt(6);fmt.line_spacing = 1.15
                fmt.keep_with_next = False
                fmt.page_break_before = False
                for run in item.runs:
                    run.font.size = Pt(14)
                if item.text.isupper():
                    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p.style.name == 'Legal Closing':
            fmt=p.paragraph_format
            fmt.left_indent=fmt.right_indent=fmt.first_line_indent=Cm(0)
            fmt.space_before=Pt(0);fmt.space_after=Pt(3);fmt.line_spacing=1.0
            fmt.keep_with_next=False
            fmt.page_break_before=False
    report['changes'].append({'action': 'reflow_text_only_layout', 'sections_before': len(sections),
                              'sections_after': len(doc.sections), 'columns_after': 1})


def reconstruct_pdf_header(doc, evidence, report):
    """Dựng hai khối đầu trang chỉ khi hình học PDF và toàn bộ chữ cùng khớp.

    Không có dữ liệu mặc định về tên cơ quan, số hiệu hoặc ngày tháng.
    """
    lines = [line for line in evidence.lines if line['page'] == 1]
    def text(line):
        chars = evidence.chars[line['start']:line['end']]
        entries = evidence.entries[line['start']:line['end']]
        return ''.join((' ' if e['space'] and i else '') + c
                       for i, (c, e) in enumerate(zip(chars, entries))).strip()
    nation = [line for line in lines if line['text'].startswith('CỘNGHÒAXÃHỘICHỦNGHĨAVIỆTNAM')]
    if len(nation) != 1:
        report['warnings'].append('Không nhận diện chắc chắn khối Quốc hiệu để dựng bố cục hai khối.')
        return
    nation = nation[0]
    agency = [line for line in lines if line['bbox'][2] < nation['bbox'][0]
              and abs(line['bbox'][1]-nation['bbox'][1]) < 5]
    motto = [line for line in lines if line['text'].startswith('Độclập-Tựdo-Hạnhphúc')
             and 0 < line['bbox'][1]-nation['bbox'][1] < 45]
    number = [line for line in lines if line['text'].startswith('Số:')
              and nation['bbox'][1] < line['bbox'][1] < nation['bbox'][1] + 100]
    dates = [line for line in lines if 'ngày' in text(line)
             and line['bbox'][0] > nation['bbox'][0]
             and nation['bbox'][1] < line['bbox'][1] < nation['bbox'][1] + 100]
    if not all(len(group) == 1 for group in (agency, motto, number, dates)):
        report['warnings'].append('Khối đầu trang không thuộc mẫu hình học đã hỗ trợ; giữ nội dung để kiểm tra.')
        return
    chunks = [text(agency[0]), text(nation), text(motto[0]), text(number[0]), text(dates[0])]
    prefix = []
    count = 0
    for p in doc.paragraphs:
        if not safe_text_paragraph(p):
            return
        if p.text.strip():
            prefix.append(p)
            count += len(compact(p.text))
        if count >= len(compact(''.join(chunks))):
            break
    if compact(''.join(p.text for p in prefix)) != compact(''.join(chunks)):
        report['warnings'].append('Chữ đầu DOCX không khớp các khối PDF; không dựng lại bằng suy đoán.')
        return
    # Giữ nguyên thứ tự chữ: cơ quan -> Quốc hiệu/Tiêu ngữ -> số hiệu/ngày.
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(5.1)
    table.columns[1].width = Cm(10.9)
    table.cell(0,0).width = table.cell(1,0).width = Cm(5.1)
    table.cell(0,1).width = table.cell(1,1).width = Cm(10.9)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        border = OxmlElement('w:' + edge);border.set(qn('w:val'), 'nil');borders.append(border)
    props.append(borders)
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    for row in table.rows:
        cant_split = OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            margins = OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'0');margin.set(qn('w:type'),'dxa');margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
    blocks = [(table.cell(0,0), [chunks[0]]), (table.cell(0,1), chunks[1:3]),
              (table.cell(1,0), [chunks[3]]), (table.cell(1,1), [chunks[4]])]
    for cell, contents in blocks:
        for i, value in enumerate(contents):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.style = 'Legal Preamble'
            fmt=p.paragraph_format
            fmt.left_indent=fmt.right_indent=fmt.first_line_indent=Cm(0)
            fmt.space_before=Pt(0);fmt.space_after=Pt(4);fmt.line_spacing=1.0
            fmt.alignment=WD_ALIGN_PARAGRAPH.CENTER
            fmt.keep_with_next=False
            run=p.add_run(value);set_font(run)
            run.font.size=Pt(14 if value == chunks[2] else 13)
            run.bold=value in chunks[:3]
            run.italic=value == chunks[4]
    prefix[0]._p.addprevious(table._tbl)
    for p in prefix:
        p._p.getparent().remove(p._p)
    report['changes'].append({'action':'reconstruct_pdf_header_two_columns',
                              'evidence':'five_text_blocks_exact_match_and_pdf_coordinates',
                              'texts':chunks})


def normalize_docx(source, destination, kind="base", pdf_evidence=None):
    doc = Document(source)
    report = {
        "rules": "nd187-2025-annexI-partial-v6",
        "legal_basis": "Phụ lục I, Mục 1, Phần I và Phần III.2 Nghị định 187/2025/NĐ-CP",
        "format_choices": {"font": "Times New Roman", "body_size_pt": 14,
                           "first_line_cm": 1.25, "paragraph_after_pt": 6,
                           "line_spacing": 1.2, "margins_mm": {"top": 20, "bottom": 20, "right": 20, "left": 30}},
        "kind": kind, "changes": [], "warnings": []}
    report["warnings"].append("Chỉ áp dụng một phần thể thức, không xác nhận văn bản đúng chuẩn hoàn toàn.")
    for number, section in enumerate(doc.sections, 1):
        if section.page_width > section.page_height:
            report["warnings"].append(f"Section {number}: giữ chiều ngang, cần xác nhận có thuộc phụ lục bảng/biểu.")
            section.page_width, section.page_height = Cm(29.7), Cm(21)
        else:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = section.right_margin = Cm(2)
        section.left_margin = Cm(3)
    before = ''.join(p.text for p in all_paragraphs(doc))
    if pdf_evidence:
        repair_from_pdf(doc, pdf_evidence, report)
    else:
        report['warnings'].append('Không có PDF: chỉ chuẩn hóa định dạng, không xác minh khoảng trắng hoặc nối đoạn.')
    subdivision_title_pending = False
    active = False
    for index, p in enumerate(list(doc.paragraphs)):
        text = p.text.strip()
        if re.match(r"^Nơi nhận\s*:", text):
            active = False
        if re.match(r"^Phụ lục(?:\s+[IVXLCDM]+)?\s*$", text, re.I):
            active = False
            report["warnings"].append(f"Đoạn {index}: phần phụ lục giữ cấu trúc, cần kiểm tra riêng.")
            break
        elif kind == "base" and (ARTICLE_START.match(text) or CHAPTER.match(text)):
            active = True
        if re.fullmatch(r"\d+", text):
            report["warnings"].append(f"Đoạn {index}: số đứng riêng {text!r}, giữ nguyên vì chưa xác minh được là số trang.")
            continue
        if not active or not text:
            if not active and text.startswith("Căn cứ "):
                # Chỉ định dạng căn cứ đã nhận diện, không tự thêm dấu hoặc sửa tên.
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Cm(1.25)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.2
                for run in p.runs:
                    run.font.size = Pt(14)
                    run.italic = True
                if "\n" in text:
                    report["warnings"].append(f"Đoạn {index}: kiểm tra nhiều căn cứ gộp trong một đoạn.")
            continue
        if not safe_text_paragraph(p):
            report["warnings"].append(f"Đoạn {index}: có thành phần phức tạp; không viết lại.")
            continue
        units = split_units(p.text)
        if len(units) > 1 or units[0] != p.text:
            paras = replace_with_units(p, units)
            report["changes"].append({"paragraph": index, "action": "normalize_or_split", "units": len(units)})
        else:
            paras = [p]
        for item in paras:
            format_unit(item, subdivision_title_pending)
            subdivision_title_pending = bool(CHAPTER.match(item.text.strip()))
            if re.search(r"(?<=[.!?])(?=[A-ZÀ-ỸĐ])", item.text):
                report["warnings"].append(f"Kiểm tra khoảng trắng sau dấu câu: {item.text[:110]}")
    # Duyệt cả bảng, header/footer để đặt font; không làm phẳng hoặc tách bảng.
    seen = set()
    for container in [doc] + [part for section in doc.sections for part in (
        section.header, section.footer, section.first_page_header,
        section.first_page_footer, section.even_page_header, section.even_page_footer)]:
        for p in all_paragraphs(container):
            if p._p in seen:
                continue
            seen.add(p._p)
            for run in p.runs:
                set_font(run)
    if pdf_evidence and kind == 'base':
        join_pdf_continuations(doc, pdf_evidence, report)
    tag_document_regions(doc, report)
    if pdf_evidence and kind == 'base':
        reflow_text_layout(doc, report)
        reconstruct_pdf_header(doc, pdf_evidence, report)
        tag_document_regions(doc, report)
    after = ''.join(p.text for p in all_paragraphs(doc))
    report['non_whitespace_length_before'] = len(compact(before))
    report['non_whitespace_length_after'] = len(compact(after))
    if compact(before) != compact(after):
        raise RuntimeError('Kiểm tra toàn vẹn thất bại: ký tự ngoài khoảng trắng đã đổi.')
    report['integrity_non_whitespace_exact_match'] = True
    report['status'] = 'needs_review'
    report['remaining_standalone_numbers'] = [
        {'output_paragraph_0based': i, 'text': p.text.strip()}
        for i, p in enumerate(doc.paragraphs) if re.fullmatch(r'\d+', p.text.strip())]
    report['layout_review'] = {
        'table_count': len(doc.tables), 'section_count': len(doc.sections),
        'requires_visual_review': True,
        'note': 'Kiểm tra khối đầu, nơi nhận/chữ ký, thứ tự cột, khoảng trống và ngắt trang; không tự dựng bảng từ chữ đã bị làm phẳng.'}
    report["warnings"].extend([
        "Không đoán cách viết đúng nếu PDF cũng dính/tách chữ hoặc không có đủ bằng chứng vị trí.",
        "Không khôi phục bảng đã mất, thứ tự nhiều cột hoặc định dạng nhấn mạnh trong đoạn được viết lại.",
        "Paragraph sau chuẩn hóa có thể khác tệp trung gian; chỉ số trong regions trỏ tới đầu ra.",
        "Chỉ dựng đầu trang khi khớp mẫu hình học và chữ PDF; chưa dựng lại khối ký/nơi nhận, số trang và đánh số riêng phụ lục.",
        "Không tự viết hoa, sửa dấu câu hay đánh lại số Điều/Khoản/Điểm; cần review chuyên môn.",
        "Tách paragraph và một khoảng trắng sau nhãn là quy tắc kỹ thuật; không phải mọi xuống dòng đều là đơn vị pháp lý.",
    ])
    doc.save(destination)
    Document(destination)  # Xác nhận DOCX mở được.
    return report


def convert_pdf_to_docx(input_path, output_path, from_docx=False, kind="base", remove_page_numbers=True):
    source, output = Path(input_path), Path(output_path)
    log = output.with_suffix('.report.json')
    if not source.is_file():
        raise FileNotFoundError(f'Không tìm thấy {source}')
    if source.suffix.lower() != ('.docx' if from_docx else '.pdf'):
        raise ValueError('Sai định dạng đầu vào; dùng --from-docx nếu chỉ thử định dạng DOCX.')
    if output.suffix.lower() != '.docx':
        raise ValueError('Đầu ra phải là .docx')
    for target in (output, log):
        if target.exists():
            raise FileExistsError(f'Đã tồn tại {target}; hãy chọn tên đầu ra mới.')
    output.parent.mkdir(parents=True, exist_ok=True)
    numbers, suspicious = (set(), []) if from_docx else inspect_pdf(source)
    with tempfile.TemporaryDirectory(prefix='pdf_to_docx_', dir=output.parent) as work:
        temporary = Path(work)
        intermediate, result = temporary / 'intermediate.docx', temporary / 'result.docx'
        margin_removals, evidence = [], None
        if from_docx:
            shutil.copy2(source, intermediate)
        else:
            from pdf2docx import Converter
            prepared_pdf = temporary / 'prepared.pdf'
            if remove_page_numbers:
                margin_removals = remove_pdf_margin_numbers(source, prepared_pdf)
            pdf_path = prepared_pdf if remove_page_numbers else source
            evidence = PdfEvidence(pdf_path)
            converter = Converter(str(pdf_path))
            try:
                converter.convert(str(intermediate))
            finally:
                converter.close()
        report = normalize_docx(intermediate, result, kind, pdf_evidence=evidence)
        report['pdf_margin_removals'] = margin_removals
        report['change_location_note'] = ('paragraph là chỉ số 0-based trong DOCX trung gian; '
                                          'output_paragraph_0based là chỉ số đầu ra. '
                                          'Sửa theo PDF có before/after và vị trí PDF.')
        report['pdf_page_number_candidates'] = sorted(numbers, key=int)
        report.update({'input': str(source.resolve()), 'output': str(output.resolve()),
                       'sparse_pdf_pages': suspicious, 'intermediate_files_retained': False})
        if suspicious:
            report['warnings'].append(f'Trang ít chữ, có thể thiếu nội dung ảnh: {suspicious}')
        if evidence and evidence.page_images:
            report['warnings'].append(f'PDF có ảnh ở các trang {evidence.page_images}; không xác minh được chữ trong ảnh, không hỗ trợ OCR.')
        report_temp = temporary / 'report.json'
        report_temp.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        # Chế độ xb tránh ghi đè tệp xuất hiện trong lúc chuyển đổi.
        created = []
        try:
            for staged, target in ((result, output), (report_temp, log)):
                with target.open('xb') as dest:
                    created.append(target)
                    with staged.open('rb') as src:
                        shutil.copyfileobj(src, dest)
        except Exception:
            for target in created:
                target.unlink(missing_ok=True)
            raise
    return output.resolve(), log.resolve()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='PDF chữ -> DOCX và report JSON; không cần DOCX đối chiếu')
    parser.add_argument('input_file')
    parser.add_argument('output_docx')
    parser.add_argument('--from-docx', action='store_true', help='Chỉ thử định dạng DOCX; không sửa theo PDF')
    parser.add_argument('--kind', choices=['base', 'other'], default='base',
                        help='base: V1, tách nhãn/nối đoạn; other: S/tham chiếu, giữ cấu trúc, sửa khoảng trắng theo PDF')
    parser.add_argument('--keep-page-numbers', action='store_true', help='Tắt loại số trang tại mép PDF')
    parser.add_argument('--remove-page-numbers', action='store_true', help='Tương thích lệnh cũ; hiện là mặc định')
    args = parser.parse_args()
    try:
        output, log = convert_pdf_to_docx(args.input_file, args.output_docx, args.from_docx,
                                         args.kind, not args.keep_page_numbers)
        print(f'DOCX: {output}\nBáo cáo: {log}')
        print('Đã tạo 2 tệp. Hãy kiểm tra các mục needs_review trong báo cáo.')
    except Exception as error:
        parser.exit(1, f'Thất bại: {error}\n')
